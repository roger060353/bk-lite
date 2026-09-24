from io import BytesIO

from django.db.models import OuterRef, Subquery
from django.http import HttpResponse
from django_filters import rest_framework as filters
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework import permissions
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.core.decorators.api_permission import HasPermission
from apps.core.utils.viewset_utils import LanguageViewSet
from apps.system_mgmt.models import Group, OpenAPICallLog
from apps.system_mgmt.serializers.openapi_call_log_serializer import OpenAPICallLogSerializer

EXPORT_HEADERS = [
    "时间",
    "源IP",
    "用户名",
    "认证时组织",
    "凭据类型",
    "钥匙ID",
    "密钥名称",
    "系统ID",
    "接口类型",
    "方法",
    "路径",
    "HTTP状态码",
    "错误码",
]


class OpenAPICallLogFilter(filters.FilterSet):
    created_at_start = filters.DateTimeFilter(field_name="created_at", lookup_expr="gte")
    created_at_end = filters.DateTimeFilter(field_name="created_at", lookup_expr="lte")
    username = filters.CharFilter(field_name="username", lookup_expr="icontains")
    credential_type = filters.CharFilter(field_name="credential_type", lookup_expr="exact")
    path = filters.CharFilter(field_name="path", lookup_expr="icontains")
    api_kind = filters.CharFilter(field_name="api_kind", lookup_expr="exact")
    outcome = filters.CharFilter(method="filter_outcome")

    class Meta:
        model = OpenAPICallLog
        fields = [
            "created_at_start",
            "created_at_end",
            "username",
            "credential_type",
            "path",
            "api_kind",
            "outcome",
        ]

    def filter_outcome(self, queryset, name, value):
        if value == "success":
            return queryset.filter(http_status__lt=400)
        if value == "failure":
            return queryset.filter(http_status__gte=400)
        return queryset


class OpenAPICallLogViewSet(LanguageViewSet):
    queryset = (
        OpenAPICallLog.objects.annotate(
            team_name=Subquery(Group.objects.filter(id=OuterRef("team_id")).values("name")[:1])
        )
        .order_by("-created_at")
    )
    serializer_class = OpenAPICallLogSerializer
    filterset_class = OpenAPICallLogFilter
    permission_classes = [permissions.IsAuthenticated]

    @HasPermission("audit_log-View")
    def list(self, request, *args, **kwargs):
        return super().list(request, *args, **kwargs)

    def get_http_method_names(self):
        if self.action == "export_excel":
            return ["post", "options"]
        return ["get", "head", "options"]

    def retrieve(self, request, *args, **kwargs):
        return Response({"result": False, "message": "Detail view is not supported"}, status=405)

    @action(detail=False, methods=["post"])
    @HasPermission("audit_log-View")
    def export_excel(self, request):
        selected_ids = request.data.get("selected_ids", [])
        queryset = self.get_queryset()
        if selected_ids:
            queryset = queryset.filter(id__in=selected_ids)
        else:
            filterset = self.filterset_class(request.data, queryset=queryset)
            if not filterset.is_valid():
                return Response({"result": False, "message": "Invalid filter parameters", "errors": filterset.errors}, status=400)
            queryset = filterset.qs

        workbook = Workbook()
        sheet = workbook.active
        sheet_title = self.loader.get("export.openapi_call_log_sheet", "API Call Log") if self.loader else "API Call Log"
        sheet.title = str(sheet_title)[:31]
        sheet.append(EXPORT_HEADERS)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center")
        for col_num, _ in enumerate(EXPORT_HEADERS, 1):
            cell = sheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        for log in queryset:
            sheet.append(
                [
                    log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else "",
                    log.source_ip,
                    log.username,
                    log.team_id if log.team_id is not None else "",
                    log.credential_type,
                    log.token_id if log.token_id is not None else "",
                    log.token_name,
                    log.system_id,
                    log.api_kind,
                    log.method,
                    log.path,
                    log.http_status,
                    log.error_code,
                ]
            )

        file_stream = BytesIO()
        workbook.save(file_stream)
        file_stream.seek(0)
        from datetime import datetime

        filename_base = self.loader.get("export.openapi_call_log_filename", "API_Call_Log") if self.loader else "API_Call_Log"
        filename = f"{filename_base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(file_stream.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
