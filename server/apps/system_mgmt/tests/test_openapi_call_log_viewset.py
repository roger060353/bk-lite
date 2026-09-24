"""OpenAPI 调用日志读侧：全量可见、筛选与 Excel 导出。"""

from datetime import timedelta
from io import BytesIO

import pytest
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.base.models import User as BaseUser
from apps.core.openapi.testing import create_api_tenant
from apps.system_mgmt.models import Group, OpenAPICallLog
from apps.system_mgmt.models import User as SystemUser

pytestmark = pytest.mark.django_db

V = "/api/v1/system_mgmt"


def _items(resp):
    data = resp.json()["data"]
    if isinstance(data, dict):
        return data.get("items") or []
    return data


def _ensure_group(team_id):
    Group.objects.get_or_create(id=team_id, defaults={"name": f"team-{team_id}"})


def _bind_system_user(user, team_id):
    SystemUser.objects.update_or_create(
        username=user.username,
        domain=user.domain,
        defaults={
            "display_name": user.username,
            "email": f"{user.username}@example.com",
            "password": "x",
            "group_list": [team_id],
        },
    )


def _client_for(user, team_id, *, is_superuser=False, permission=None):
    user.is_superuser = is_superuser
    user.save()
    user.group_list = [{"id": team_id}]
    user.permission = permission or {"system-manager": {"audit_log-View"}}
    client = APIClient()
    client.force_authenticate(user=user)
    client.cookies["current_team"] = str(team_id)
    return client


def _log(**over):
    defaults = {
        "source_ip": "10.0.0.1",
        "username": "alice",
        "credential_type": "api_token",
        "token_name": "alice-key",
        "system_id": "",
        "team_id": 1,
        "api_kind": OpenAPICallLog.API_KIND_INTERNAL,
        "method": "GET",
        "path": "/openapi/v1/patch-mgmt/module-data",
        "http_status": 200,
        "error_code": "",
    }
    defaults.update(over)
    return OpenAPICallLog.objects.create(**defaults)


def test_auditor_sees_every_org_and_empty_username():
    _ensure_group(1)
    _ensure_group(2)
    user_a, _ = create_api_tenant(1, username="alice-org")
    user_b, _ = create_api_tenant(2, username="bob-org")
    _bind_system_user(user_a, 1)
    _bind_system_user(user_b, 2)
    own = _log(username=user_a.username, team_id=1)
    other = _log(username=user_b.username, team_id=2, path="/openapi/v1/cmdb/instances")
    empty = _log(username="", credential_type="", team_id=None, path="/openapi/v1/unknown")

    auditor = BaseUser.objects.create_user(username="auditor", password="pw", domain="domain.com")
    _bind_system_user(auditor, 1)
    client = _client_for(auditor, 1)
    resp = client.get(f"{V}/openapi_call_log/")
    assert resp.status_code == 200
    items = _items(resp)
    ids = {item["id"] for item in items}
    assert own.id in ids
    assert other.id in ids
    assert empty.id in ids
    by_id = {item["id"]: item for item in items}
    assert by_id[own.id]["team_name"] == Group.objects.get(id=1).name
    assert by_id[other.id]["team_name"] == Group.objects.get(id=2).name
    assert by_id[empty.id]["team_name"] is None
    assert "token_id" not in by_id[own.id]
    detail = client.get(f"{V}/openapi_call_log/{own.id}/")
    assert detail.status_code == 405


def test_toolbar_filters_hit_username_token_kind_and_outcome():
    _ensure_group(1)
    user_a, _ = create_api_tenant(1, username="filter-user")
    _bind_system_user(user_a, 1)
    _log(username=user_a.username, team_id=1, token_name="other-key")
    target = _log(
        username=user_a.username,
        team_id=1,
        api_kind=OpenAPICallLog.API_KIND_EXTERNAL,
        method="PUT",
        path="/openapi/v1/itsm/tickets/99",
        http_status=403,
        error_code="SCOPE_DENIED",
        credential_type="system_token",
        token_name="itsm-bot",
        system_id="itsm",
    )
    now = timezone.now()
    admin = BaseUser.objects.create_user(username="filteradmin", password="pw", domain="domain.com")
    client = _client_for(admin, 1, is_superuser=True)
    resp = client.get(
        f"{V}/openapi_call_log/",
        {
            "page_size": 20,
            "created_at_start": (now - timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S"),
            "created_at_end": (now + timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S"),
            "username": user_a.username,
            "credential_type": "system_token",
            "path": "tickets",
            "api_kind": "external",
            "outcome": "failure",
        },
    )
    assert resp.status_code == 200
    items = _items(resp)
    assert [item["id"] for item in items] == [target.id]

    success = client.get(
        f"{V}/openapi_call_log/",
        {"username": user_a.username, "outcome": "success"},
    )
    assert [item["id"] for item in _items(success)] != [target.id]
    assert target.id not in {item["id"] for item in _items(success)}


def test_export_excel_selected_and_filtered_has_no_payload():
    _ensure_group(1)
    user_a, _ = create_api_tenant(1, username="export-user")
    _bind_system_user(user_a, 1)
    row = _log(username=user_a.username, team_id=1, token_name="export-key")
    admin = BaseUser.objects.create_user(username="exportadmin", password="pw", domain="domain.com")
    client = _client_for(admin, 1, is_superuser=True)
    resp = client.post(f"{V}/openapi_call_log/export_excel/", {"selected_ids": [row.id]}, format="json")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp["Content-Type"]
    workbook = load_workbook(BytesIO(resp.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == [
        "时间",
        "源IP",
        "用户名",
        "认证时组织",
        "凭据类型",
        "钥匙ID",
        "密钥名称",
        "系统ID",
        "接口类型",
        "方法",
        "路径",
        "HTTP状态码",
        "错误码",
    ]
    assert "服务名" not in headers
    assert "请求正文" not in "".join(str(h) for h in headers)
    assert "响应正文" not in "".join(str(h) for h in headers)
    body = " ".join(str(cell.value) for row_cells in sheet.iter_rows(min_row=2) for cell in row_cells)
    assert "export-key" in body
    assert user_a.username in body

    filtered = client.post(
        f"{V}/openapi_call_log/export_excel/",
        {"api_kind": "internal", "credential_type": "api_token"},
        format="json",
    )
    assert filtered.status_code == 200
    filtered_book = load_workbook(BytesIO(filtered.content))
    assert filtered_book.active.max_row >= 2


def test_system_id_is_required_only_for_system_token():
    _log(credential_type="system_token", system_id="itsm", token_name="itsm-key")
    _log(credential_type="api_token", system_id="", token_name="personal")
    _log(credential_type="", system_id="", username="", token_name="")
    with pytest.raises(IntegrityError):
        with transaction.atomic():
            _log(credential_type="system_token", system_id="")
    with pytest.raises(IntegrityError):
        with transaction.atomic():
            _log(credential_type="api_token", system_id="itsm")
