"""实例导出的文件内容、权限范围和图库访问预算。"""

import json
import logging
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest

from apps.cmdb.services.instance import InstanceManage
from apps.cmdb.services.model import ModelManage


def uuid(number):
    return f"123e4567-e89b-42d3-a456-{number:012d}"


ATTRS = [{"attr_id": "inst_name", "attr_name": "实例名", "attr_type": "str"}]
ASSOCIATION = {
    "model_asst_id": "host_belong_app",
    "asst_id": "belong",
    "src_model_id": "host",
    "dst_model_id": "app",
    "src_model_name": "主机",
    "dst_model_name": "应用",
}


def read_rows(stream):
    return list(openpyxl.load_workbook(stream).active.values)


def test_association_export_batches_queries_and_preserves_cells(monkeypatch, fake_graph):
    instances = [dict(_id=i, inst_uuid=uuid(i), inst_name=f"host-{i}") for i in range(1, 101)]
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    monkeypatch.setattr(ModelManage, "model_association_search", lambda *a, **k: [ASSOCIATION])
    monkeypatch.setattr(
        ModelManage,
        "search_model",
        lambda: [
            {"model_id": "host", "model_name": "主机"},
            {"model_id": "app", "model_name": "应用"},
        ],
    )
    graph = fake_graph(
        "apps.cmdb.services.instance",
        query_entity=(instances, None),
        query_entity_by_id=lambda i: instances[i - 1],
        query_edge=[],
        query_export_associations=[
            {
                "inst_uuid": uuid(1),
                "model_asst_id": "host_belong_app",
                "peer_uuid": uuid(200),
                "peer_name": "应用一",
                "edge_id": 1,
                "src_model_id": "host",
                "dst_model_id": "app",
            }
        ],
    )
    stream = InstanceManage.inst_export(
        "host",
        ids=[],
        permissions_map={1: {"inst_names": []}},
        attr_list=["inst_name"],
        association_list=["host_belong_app"],
    )
    rows = read_rows(stream)
    assert rows[2] == ("字段标识(请勿编辑)", "inst_name", "host_belong_app")
    assert rows[3] == (None, "host-1", "应用一")
    assert rows[4] == (None, "host-2", None)
    assert len(rows) == 103
    assert not [call for call in graph.calls if call[0] in {"query_entity_by_id", "query_edge"}]
    assert len([call for call in graph.calls if call[0] == "query_export_associations"]) == 1


def test_unselected_user_and_organization_fields_do_not_load_options(monkeypatch):
    monkeypatch.setattr(
        ModelManage,
        "search_model_info",
        lambda _: {
            "attrs": json.dumps(
                ATTRS
                + [
                    {"attr_id": "operator", "attr_name": "维护人", "attr_type": "user"},
                    {"attr_id": "organization", "attr_name": "组织", "attr_type": "organization"},
                ]
            )
        },
    )
    monkeypatch.setattr("apps.cmdb.services.model.build_unique_rule_context", lambda _: SimpleNamespace(unique_rules=[]))
    monkeypatch.setattr("apps.cmdb.language.service.apply_attr_translations", lambda attrs, *a: attrs)
    users = Mock(side_effect=AssertionError("不应加载未选中的用户选项"))
    groups = Mock(side_effect=AssertionError("不应加载未选中的组织选项"))
    monkeypatch.setattr(ModelManage, "get_cached_user_options", users)
    monkeypatch.setattr(ModelManage, "get_cached_organization_options", groups)
    attrs = ModelManage.search_model_attr_v2("host", attr_ids=["inst_name"])
    assert [attr["attr_id"] for attr in attrs] == ["inst_name"]
    users.assert_not_called()
    groups.assert_not_called()


def test_export_without_associations_skips_association_metadata(monkeypatch, fake_graph):
    metadata = Mock(return_value=ATTRS)
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", metadata)
    associations = Mock(side_effect=AssertionError("不应加载未选中的关联"))
    monkeypatch.setattr(ModelManage, "model_association_search", associations)
    graph = fake_graph("apps.cmdb.services.instance", query_entity=([{"_id": 1, "inst_name": "host-1"}], None))
    rows = read_rows(InstanceManage.inst_export("host", [], {1: {"inst_names": []}}, attr_list=["inst_name"]))
    assert rows[3] == (None, "host-1")
    metadata.assert_called_once_with("host", attr_ids=["inst_name"])
    associations.assert_not_called()
    assert not [call for call in graph.calls if call[0] == "query_export_associations"]


def test_export_logs_only_bounded_completion_and_preserves_failure(monkeypatch, fake_graph, caplog):
    sentinel = "export-sensitive-payload-7651"
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    fake_graph("apps.cmdb.services.instance", query_entity=([{"_id": 1, "inst_name": sentinel}], None))
    with caplog.at_level(logging.INFO, logger="cmdb"):
        rows = read_rows(InstanceManage.inst_export(sentinel, [1], {1: {"inst_names": []}}))
    assert rows[3][1] == sentinel
    records = [record for record in caplog.records if record.name == "cmdb"]
    assert len(records) == 1
    record = records[0]
    assert record.msg.startswith("event=cmdb_instance_exported export_id=%s")
    assert record.args and "rows=1" in record.getMessage()
    assert sentinel not in "\n".join(logging.Formatter().format(item) + repr(item.args) for item in records)
    assert record.exc_info is None
    caplog.clear()
    error = RuntimeError(sentinel)

    def fail(**kwargs):
        raise error

    fake_graph("apps.cmdb.services.instance", query_entity=fail)
    with pytest.raises(RuntimeError) as caught:
        InstanceManage.inst_export("host", [], {1: {"inst_names": []}})
    assert caught.value is error
    # 此服务继续上抛；不抢上层异常处理边界的 traceback 所有权。
    assert not [record for record in caplog.records if record.name == "cmdb"]


@pytest.mark.parametrize("team", [1, 2])
def test_export_pages_by_uuid_preserving_organization_scope(monkeypatch, fake_graph, team):
    from apps.cmdb.graph.export_query import EXPORT_BATCH_SIZE

    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    calls = []
    data = [dict(_id=i, inst_uuid=uuid(i), inst_name=f"host-{i}", organization=[1]) for i in range(1, 502)]
    data.append(dict(_id=600, inst_uuid=uuid(600), inst_name="other-team", organization=[2]))

    def query(**kwargs):
        calls.append(kwargs)
        assert kwargs["page"] == {"skip": 0, "limit": EXPORT_BATCH_SIZE}
        assert kwargs["include_count"] is False
        assert kwargs["order"] == "inst_uuid"
        assert set(kwargs["fields"]) == {"inst_name", "inst_uuid", "model_id"}
        assert kwargs["format_permission_dict"] == {team: []}
        cursor = next((p["value"] for p in kwargs["params"] if p["type"] == "str>"), "")
        # 模拟上一批返回后删除已导出的行；keyset 不会像 offset 那样跳过下一批。
        if cursor:
            data[:] = [item for item in data if item["inst_uuid"] > cursor]
        return [item for item in data if team in item["organization"] and item["inst_uuid"] > cursor][:EXPORT_BATCH_SIZE], None

    fake_graph("apps.cmdb.services.instance", query_entity=query)
    rows = read_rows(InstanceManage.inst_export("host", [], {team: {"inst_names": []}}, attr_list=["inst_name"]))
    expected = [f"host-{i}" for i in range(1, 502)] if team == 1 else ["other-team"]
    assert [row[1] for row in rows[3:]] == expected
    assert len(calls) == (2 if team == 1 else 1)


@pytest.mark.parametrize("count", [1, 500, 501, 1001])
def test_empty_associations_cost_per_batch_not_per_instance(monkeypatch, fake_graph, count):
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    monkeypatch.setattr(ModelManage, "model_association_search", lambda *a, **k: [ASSOCIATION])

    def query(**kwargs):
        cursor = next((p["value"] for p in kwargs["params"] if p["type"] == "str>"), "")
        return [dict(_id=i, inst_uuid=uuid(i), inst_name=f"h{i}") for i in range(1, count + 1) if uuid(i) > cursor][:500], None

    graph = fake_graph("apps.cmdb.services.instance", query_entity=query, query_export_associations=[])
    rows = read_rows(InstanceManage.inst_export("host", [], {1: {"inst_names": []}}, association_list=["host_belong_app"]))
    assert len(rows) == count + 3
    assert all(row[-1] is None for row in rows[3:])
    assert len([call for call in graph.calls if call[0] == "query_export_associations"]) == (count + 499) // 500
    assert not [call for call in graph.calls if call[0] in {"query_edge", "query_entity_by_id"}]


def test_relation_cells_preserve_directions_and_distinct_same_name_edges(monkeypatch, fake_graph):
    reverse = dict(ASSOCIATION, model_asst_id="app_belong_host", src_model_id="app", dst_model_id="host")
    self_relation = dict(ASSOCIATION, model_asst_id="host_belong_host", dst_model_id="host")
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    monkeypatch.setattr(ModelManage, "model_association_search", lambda *a, **k: [ASSOCIATION, reverse, self_relation])

    def relation(edge_id, definition, name):
        return dict(
            inst_uuid=uuid(1),
            model_asst_id=definition["model_asst_id"],
            edge_id=edge_id,
            peer_name=name,
            src_model_id=definition["src_model_id"],
            dst_model_id=definition["dst_model_id"],
        )

    loop = relation(4, self_relation, "self")
    unrelated = dict(relation(5, ASSOCIATION, "hidden"), model_asst_id="host_belong_hidden")
    wrong_model = dict(relation(6, ASSOCIATION, "wrong"), dst_model_id="hidden")
    graph = fake_graph(
        "apps.cmdb.services.instance",
        query_entity=([dict(_id=1, inst_uuid=uuid(1), inst_name="h1")], None),
        query_export_associations=[
            relation(1, ASSOCIATION, "same"),
            relation(2, ASSOCIATION, "same"),
            relation(3, reverse, "incoming"),
            loop,
            loop,
            unrelated,
            wrong_model,
        ],
    )
    rows = read_rows(
        InstanceManage.inst_export("host", [1], {1: {"inst_names": []}}, association_list=["host_belong_app", "app_belong_host", "host_belong_host"])
    )
    assert rows[3] == (None, "h1", "same,same", "incoming", "self")
    query = next(call for call in graph.calls if call[0] == "query_export_associations")
    assert query[1][1] == [uuid(1)]


def test_empty_permission_scope_never_reads_instances(monkeypatch, fake_graph):
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    graph = fake_graph("apps.cmdb.services.instance")
    assert len(read_rows(InstanceManage.inst_export("host", [], {}))) == 3
    assert graph.calls == []


def test_failed_later_batch_cleans_temporary_xml_and_preserves_exception(monkeypatch, fake_graph):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    calls = 0
    error = RuntimeError("later batch failed")

    def query(**kwargs):
        nonlocal calls
        calls += 1
        if calls > 1:
            raise error
        return [dict(_id=i, inst_uuid=uuid(i), inst_name=f"h{i}") for i in range(1, 501)], None

    fake_graph("apps.cmdb.services.instance", query_entity=query)
    before = set(ALL_TEMP_FILES)
    with pytest.raises(RuntimeError) as caught:
        InstanceManage.inst_export("host", [], {1: {"inst_names": []}}, file_backed=True)
    assert caught.value is error
    assert set(ALL_TEMP_FILES) == before


def test_file_backed_export_returns_seekable_file(monkeypatch, fake_graph):
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    fake_graph("apps.cmdb.services.instance", query_entity=([dict(_id=1, inst_name="host-1")], None))
    with InstanceManage.inst_export("host", [1], {1: {"inst_names": []}}, file_backed=True) as stream:
        assert stream.fileno() >= 0
        assert stream.tell() == 0
        assert read_rows(stream)[3] == (None, "host-1")


def test_selected_uuid_resolution_fetches_only_identity(monkeypatch, fake_graph):
    graph = fake_graph(
        "apps.cmdb.services.instance",
        query_entity=(
            [
                {"_id": 1, "inst_uuid": uuid(1), "model_id": "host"},
                {"_id": 2, "inst_uuid": uuid(2), "model_id": "host"},
            ],
            None,
        ),
    )
    items = InstanceManage.query_entity_by_uuids([uuid(2), uuid(1)], fields=["inst_uuid", "model_id"])
    assert [item["_id"] for item in items] == [2, 1]
    assert len(graph.calls) == 1
    assert graph.calls[0][0] == "query_entity"
    assert graph.calls[0][2]["fields"] == ["inst_uuid", "model_id"]


def test_file_save_failure_closes_file_without_replacing_error(monkeypatch, fake_graph):
    from io import BytesIO

    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    fake_graph("apps.cmdb.services.instance", query_entity=([], None))
    stream = BytesIO()
    monkeypatch.setattr("apps.cmdb.services.instance.TemporaryFile", lambda **kwargs: stream)
    error = OSError("disk full")

    def fail_save(self, target):
        raise error

    monkeypatch.setattr(openpyxl.Workbook, "save", fail_save)
    before = set(ALL_TEMP_FILES)
    with pytest.raises(OSError) as caught:
        InstanceManage.inst_export("host", [], {1: {"inst_names": []}}, file_backed=True)
    assert caught.value is error
    assert stream.closed
    assert set(ALL_TEMP_FILES) == before


def test_async_export_budget_and_progress(monkeypatch, fake_graph):
    from apps.cmdb.services.transfer_service import TransferError

    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    fake_graph("apps.cmdb.services.instance", query_entity=([{"inst_uuid": uuid(i), "inst_name": str(i)} for i in range(3)], None))
    progress = []
    with pytest.raises(TransferError) as error:
        InstanceManage.inst_export(
            "host", [], {1: {"inst_names": []}}, attr_list=["inst_name"], row_limit=2, progress=lambda count: progress.append(count)
        )
    assert error.value.code == "export_limit"
    assert progress == [0]


def test_export_keeps_user_text_as_text_instead_of_executable_formula(monkeypatch, fake_graph):
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    fake_graph("apps.cmdb.services.instance", query_entity=([{"inst_name": '=HYPERLINK("private")'}], None))
    stream = InstanceManage.inst_export("host", [], {1: {"inst_names": []}}, attr_list=["inst_name"])
    cell = openpyxl.load_workbook(stream).active.cell(4, 2)
    assert cell.value == '=HYPERLINK("private")'
    assert cell.data_type == "s"
