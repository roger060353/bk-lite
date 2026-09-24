import io

import openpyxl
import pytest

from apps.cmdb.services.transfer_service import TransferError
from apps.cmdb.services.transfer_validation import inspect_workbook


def workbook(model="host", rows=1):
    book = openpyxl.Workbook()
    book.active.title = model
    for row in (["实例名"], ["属性"], ["inst_name"]):
        book.active.append(row)
    for _ in range(rows):
        book.active.append(["host-1"])
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return output


def test_inspect_workbook_rejects_wrong_model_and_counts_actual_rows():
    result = inspect_workbook(workbook(rows=3), "host", allowed_fields={"inst_name"})
    assert result["rows"] == 3
    assert len(result["sha256"]) == 64
    with pytest.raises(TransferError) as error:
        inspect_workbook(workbook(model="mysql"), "host", allowed_fields={"inst_name"})
    assert error.value.code == "invalid_workbook"
    with pytest.raises(TransferError):
        inspect_workbook(workbook(), "host", allowed_fields={"organization"})


def test_accepts_export_template_marker_column():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "host"
    sheet.append(["字段名", "实例名"])
    sheet.append(["字段类型", "str"])
    sheet.append(["字段标识(请勿编辑)", "inst_name"])
    sheet.append([None, "host-one"])
    stream = io.BytesIO()
    book.save(stream)
    assert inspect_workbook(stream, "host", allowed_fields={"inst_name"})["rows"] == 1


@pytest.mark.parametrize("kind", ["formula", "too_many_rows", "too_many_columns", "duplicate_header", "empty", "invalid_zip"])
def test_rejects_unsafe_or_oversized_workbooks(kind):
    stream = workbook()
    if kind == "invalid_zip":
        stream = io.BytesIO(b"not-a-workbook")
    else:
        book = openpyxl.load_workbook(stream)
        sheet = book.active
        if kind == "formula":
            sheet.cell(4, 1, "=1+1")
        elif kind == "too_many_rows":
            sheet.cell(10004, 1, "bad")
        elif kind == "too_many_columns":
            sheet.cell(1, 257, "bad")
        elif kind == "duplicate_header":
            sheet.cell(3, 2, "inst_name")
        else:
            sheet.delete_rows(4)
        stream = io.BytesIO()
        book.save(stream)
    with pytest.raises(TransferError):
        inspect_workbook(stream, "host", allowed_fields={"inst_name"})


@pytest.mark.parametrize(
    "params",
    [
        {"scope": "all", "inst_uuids": ["123e4567-e89b-42d3-a456-426614174000"]},
        {"scope": "selected", "inst_uuids": []},
        {"scope": "all", "attr_list": ["inst_name", "inst_name"]},
        {"scope": "all", "association_list": ["a", "a"]},
    ],
)
def test_export_request_refuses_ambiguous_scope(params):
    from apps.cmdb.services.transfer_validation import ExportRequest

    serializer = ExportRequest(data={"model_id": "host", "attr_list": ["inst_name"], **params})
    assert not serializer.is_valid()
