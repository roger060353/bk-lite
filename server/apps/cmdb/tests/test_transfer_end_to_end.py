"""真实 ORM/工作簿/HTTP/执行链路；只替换外部身份、图库元数据与对象存储。"""

import io
import json
from contextlib import contextmanager
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.cmdb.services.model import ModelManage
from apps.cmdb.services.transfer_authorization import TransferAuthorization
from apps.cmdb.services.transfer_execution import TransferExecution
from apps.cmdb.services.transfer_service import TransferService
from apps.cmdb.views.transfer_task import TransferTaskViewSet

ATTRS = [{"attr_id": "inst_name", "attr_name": "实例名", "attr_type": "str", "is_only": True, "is_required": True, "editable": True}]
UUID = "123e4567-e89b-42d3-a456-426614174000"
INSTANCE = {"_id": 1, "inst_uuid": UUID, "inst_name": "host-one", "model_id": "host", "organization": [1]}


class MemoryFiles:
    MAX_BYTES = 100 * 1024 * 1024

    def __init__(self):
        self.objects = {}

    def put(self, key, stream):
        stream.seek(0)
        self.objects[key] = stream.read()
        return {"key": key, "size": len(self.objects[key])}

    @contextmanager
    def local_copy(self, key, **kwargs):
        with io.BytesIO(self.objects[key]) as stream:
            yield stream

    read = local_copy


@pytest.fixture
def authorized(monkeypatch):
    monkeypatch.setattr(
        "apps.cmdb.services.transfer_authorization.build_user_authorization_context",
        lambda transfer_owner: {
            "username": transfer_owner.username,
            "domain": transfer_owner.domain,
            "is_superuser": True,
            "permission": {},
            "group_list": [{"id": 1}],
            "roles": ["admin"],
        },
    )
    monkeypatch.setattr("apps.cmdb.utils.permission_util.get_permission_rules", lambda **kw: {"team": [1]})
    monkeypatch.setattr(
        "apps.cmdb.services.model_visibility.BusinessModelVisibility.resolve", lambda ids: {"host": {"model_id": "host", "model_name": "主机"}}
    )
    monkeypatch.setattr(ModelManage, "search_model_attr", lambda *a, **k: ATTRS)
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: ATTRS)
    monkeypatch.setattr(ModelManage, "model_association_search", lambda *a, **k: [])


def create_request(transfer_owner, data, key="request"):
    request = APIRequestFactory().post("/cmdb/api/transfer_tasks/export/", data, format="json", HTTP_IDEMPOTENCY_KEY=key)
    request.COOKIES["current_team"] = "1"
    force_authenticate(request, user=SimpleNamespace(username=transfer_owner.username, domain=transfer_owner.domain, is_authenticated=True))
    return TransferTaskViewSet.as_view({"post": "export_file"})(request)


@pytest.mark.parametrize("scope", ["all", "selected", "currentPage"])
def test_export_http_to_readable_file_and_download_rechecks_resource_owner(transfer_owner, authorized, fake_graph, monkeypatch, scope):
    files = MemoryFiles()
    fake_graph("apps.cmdb.services.instance", query_entity=([dict(INSTANCE)], None))
    from apps.cmdb.services.instance import InstanceManage

    current = [dict(INSTANCE)]
    monkeypatch.setattr(InstanceManage, "query_entity_by_uuids", lambda ids, **kw: current)
    response = create_request(
        transfer_owner, {"model_id": "host", "scope": scope, "attr_list": ["inst_name"], "inst_uuids": [] if scope == "all" else [UUID]}
    )
    assert response.status_code == 202
    payload = json.loads(response.content)["data"]
    assert payload["status"] == "queued"
    TransferExecution.run(payload["task_id"], files=files)
    task = TransferService.get(transfer_owner, payload["task_id"])
    assert task.status == "succeeded"
    assert task.summary == {"exported": 1}
    sheet = openpyxl.load_workbook(io.BytesIO(files.objects[task.artifacts["result"]["key"]])).active
    assert sheet.cell(4, 2).value == "host-one"
    TransferExecution.validate_download(task, files)
    current[0]["organization"] = [99]
    from apps.cmdb.services.transfer_service import TransferError

    with pytest.raises(TransferError) as error:
        TransferExecution.validate_download(task, files)
    assert error.value.status_code == 403


@pytest.mark.parametrize("update_existing", [False, True])
def test_import_partial_success_keeps_operation_audit_and_row_numbers(transfer_owner, authorized, fake_graph, monkeypatch, update_existing):
    from apps.cmdb.models.operation import CmdbOperation, CmdbOperationOutbox
    from apps.cmdb.services.instance import InstanceManage
    from apps.cmdb.services.transfer_validation import inspect_workbook

    # 图库写边界替身；接纳、校验、计数、操作账本和审计 Outbox 使用真实实现。
    writes = Mock(side_effect=lambda model, data, operator, **kwargs: dict(INSTANCE, **data))
    monkeypatch.setattr(InstanceManage, "instance_create", writes)
    updates = Mock(side_effect=lambda teams, roles, uuid, data, operator, **kwargs: dict(INSTANCE, **data))
    monkeypatch.setattr(InstanceManage, "instance_update_by_uuid", updates)
    monkeypatch.setattr(
        "apps.cmdb.utils.Import.build_unique_rule_context", lambda _: SimpleNamespace(unique_rules=[], attrs_by_id={"inst_name": ATTRS[0]})
    )
    fake_graph("apps.cmdb.services.transfer_import", query_entity=([dict(INSTANCE, inst_name="one")] if update_existing else [], 0))
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "host"
    for row in (["实例名"], ["str"], ["inst_name"], ["one"], ["one"], ["two"]):
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    info = inspect_workbook(stream, "host", allowed_fields={"inst_name"})
    files = MemoryFiles()
    source = "transfer/tmp/test/source.xlsx"
    files.put(source, stream)
    context = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "import")
    task = TransferService.submit(
        owner=transfer_owner,
        kind="import",
        model_id="host",
        team_id=1,
        include_children=False,
        params={},
        authorization=context.snapshot,
        schema_hash=context.schema_hash,
        idempotency_key="import",
        source_key=source,
        source_hash=info["sha256"],
    )
    TransferExecution.run(task.pk, files=files)
    task = TransferService.get(transfer_owner, task.pk)
    assert task.status == "partial_success"
    assert task.summary == {
        "created": 1 if update_existing else 2,
        "updated": 1 if update_existing else 0,
        "failed_rows": 1,
        "created_relations": 0,
        "failed_relations": 0,
    }
    assert writes.call_count + updates.call_count == 2
    assert CmdbOperation.objects.filter(idempotency_key__startswith=f"transfer:{task.pk}:").count() == 2
    assert CmdbOperationOutbox.objects.filter(event_type="change_record").count() == 2
    report = openpyxl.load_workbook(io.BytesIO(files.objects[task.artifacts["errors"]["key"]])).active
    assert report.cell(2, 1).value == 5
    assert report.cell(2, 2).value == "instance"


@pytest.mark.parametrize("source_changed", [False, True])
def test_upload_replay_and_download_http_use_owner_scope(transfer_owner, authorized, monkeypatch, source_changed):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.cmdb.tests.test_transfer_validation import workbook

    files = MemoryFiles()
    monkeypatch.setattr("apps.cmdb.views.transfer_task.TransferFiles", lambda: files)

    def upload(key):
        file = SimpleUploadedFile("host.xlsx", workbook().getvalue())
        request = APIRequestFactory().post("/", {"model_id": "host", "file": file}, format="multipart", HTTP_IDEMPOTENCY_KEY=key)
        request.COOKIES["current_team"] = "1"
        force_authenticate(request, user=SimpleNamespace(username=transfer_owner.username, domain=transfer_owner.domain, is_authenticated=True))
        return TransferTaskViewSet.as_view({"post": "import_file"})(request)

    response = upload("one")
    assert response.status_code == 202
    task_id = json.loads(response.content)["data"]["task_id"]
    assert json.loads(upload("one").content)["data"]["task_id"] == task_id
    assert TransferService.list(transfer_owner).count() == 1
    if source_changed:
        task = TransferService.get(transfer_owner, task_id)
        files.objects[task.source_key] = workbook(rows=2).getvalue()
        TransferExecution.run(task_id, files=files)
        assert TransferService.get(transfer_owner, task_id).error_code == "source_changed"
        return
    token = TransferService.claim(task_id)
    task = TransferService.get(transfer_owner, task_id)
    report = files.put(f"transfer/{transfer_owner.pk}/{task.pk}/{token}/errors.xlsx", io.BytesIO(b"error-report"))
    TransferService.finish(task_id, token, "partial_success", artifacts={"errors": report})
    from apps.cmdb.tests.test_transfer_views import request_view

    downloaded = request_view(transfer_owner, "download", "post", task_id, {"artifact": "errors"})
    assert downloaded.status_code == 200
    assert b"".join(downloaded.streaming_content) == b"error-report"
    downloaded.close()
    assert request_view(transfer_owner, "download", "post", task_id, {"artifact": "manifest"}).status_code == 404


def test_storage_failure_does_not_evict_existing_history(transfer_owner, authorized, monkeypatch):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.cmdb.tests.test_transfer_service import submit
    from apps.cmdb.tests.test_transfer_validation import workbook

    for index in range(5):
        task = submit(transfer_owner, key=str(index))
        TransferService.cancel(transfer_owner, task.pk)
    history = list(TransferService.list(transfer_owner).values_list("pk", flat=True))
    files = MemoryFiles()
    files.put = Mock(side_effect=OSError("storage unavailable"))
    monkeypatch.setattr("apps.cmdb.views.transfer_task.TransferFiles", lambda: files)
    request = APIRequestFactory().post(
        "/", {"model_id": "host", "file": SimpleUploadedFile("host.xlsx", workbook().getvalue())}, format="multipart", HTTP_IDEMPOTENCY_KEY="upload"
    )
    request.COOKIES["current_team"] = "1"
    force_authenticate(request, user=SimpleNamespace(username=transfer_owner.username, domain=transfer_owner.domain, is_authenticated=True))
    response = TransferTaskViewSet.as_view({"post": "import_file"})(request)
    assert response.status_code == 503
    assert list(TransferService.list(transfer_owner).values_list("pk", flat=True)) == history


def test_failed_export_retry_replaces_old_record_and_replays_same_request(transfer_owner, authorized):
    payload = {"model_id": "host", "scope": "all", "attr_list": ["inst_name"]}
    first = json.loads(create_request(transfer_owner, payload).content)["data"]["task_id"]
    token = TransferService.claim(first)
    TransferService.finish(first, token, "failed", code="export_limit")
    request = APIRequestFactory().post("/", {}, format="json", HTTP_IDEMPOTENCY_KEY="retry")
    request.COOKIES["current_team"] = "1"
    force_authenticate(request, user=SimpleNamespace(username=transfer_owner.username, domain=transfer_owner.domain, is_authenticated=True))
    response = TransferTaskViewSet.as_view({"post": "retry"})(request, pk=first)
    assert response.status_code == 202
    replacement = json.loads(response.content)["data"]["task_id"]
    assert replacement != first
    assert [str(task.pk) for task in TransferService.list(transfer_owner)] == [replacement]
    from apps.cmdb.tests.test_transfer_views import request_view

    assert request_view(transfer_owner, "retrieve", task_id=first).status_code == 404
    repeated = APIRequestFactory().post("/", {}, format="json", HTTP_IDEMPOTENCY_KEY="retry")
    force_authenticate(repeated, user=SimpleNamespace(username=transfer_owner.username, domain=transfer_owner.domain, is_authenticated=True))
    replay = TransferTaskViewSet.as_view({"post": "retry"})(repeated, pk=first.upper())
    assert replay.status_code == 202
    assert json.loads(replay.content)["data"]["task_id"] == replacement


def test_retry_validation_failure_keeps_old_failed_record(transfer_owner, authorized, monkeypatch):
    payload = {"model_id": "host", "scope": "all", "attr_list": ["inst_name"]}
    first = json.loads(create_request(transfer_owner, payload).content)["data"]["task_id"]
    token = TransferService.claim(first)
    TransferService.finish(first, token, "failed", code="export_limit")
    monkeypatch.setattr(ModelManage, "search_model_attr", lambda *args, **kwargs: [])
    request = APIRequestFactory().post("/", {}, format="json", HTTP_IDEMPOTENCY_KEY="retry")
    request.COOKIES["current_team"] = "1"
    force_authenticate(request, user=SimpleNamespace(username=transfer_owner.username, domain=transfer_owner.domain, is_authenticated=True))
    response = TransferTaskViewSet.as_view({"post": "retry"})(request, pk=first)
    assert response.status_code == 400
    assert [str(task.pk) for task in TransferService.list(transfer_owner)] == [first]
    assert TransferService.get(transfer_owner, first).status == "failed"


@pytest.mark.parametrize("existing_relation", [False, True])
def test_relation_failures_are_separate_from_successful_instance_rows(transfer_owner, authorized, fake_graph, monkeypatch, existing_relation):
    from apps.cmdb.services.instance import InstanceManage
    from apps.cmdb.services.transfer_validation import inspect_workbook

    association = {"model_asst_id": "host_belong_app", "asst_id": "belong", "src_model_id": "host", "dst_model_id": "app"}
    monkeypatch.setattr(ModelManage, "model_association_search", lambda *a, **k: [association])
    monkeypatch.setattr("apps.cmdb.services.model_visibility.BusinessModelVisibility.resolve", lambda ids: {key: {"model_id": key} for key in ids})
    monkeypatch.setattr(
        "apps.cmdb.utils.Import.build_unique_rule_context", lambda _: SimpleNamespace(unique_rules=[], attrs_by_id={"inst_name": ATTRS[0]})
    )
    monkeypatch.setattr(InstanceManage, "instance_create", lambda model, data, operator, **kw: dict(INSTANCE, **data))
    from apps.core.exceptions.base_app_exception import BaseAppException

    edges = Mock(return_value={}, side_effect=BaseAppException("instance association repetition") if existing_relation else None)
    monkeypatch.setattr(InstanceManage, "instance_association_create_by_uuid", edges)

    def query(label, params, **kwargs):
        if params[0]["value"] == "app" and params[1]["value"] == "peer":
            return [dict(INSTANCE, inst_name="peer", model_id="app", inst_uuid="123e4567-e89b-42d3-a456-426614174001")], 1
        return [], 0

    fake_graph("apps.cmdb.services.transfer_import", query_entity=query)
    book = openpyxl.Workbook()
    book.active.title = "host"
    for row in (["实例名", "关联"], ["str", "关联"], ["inst_name", "host_belong_app"], ["one", "peer,missing"]):
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    info = inspect_workbook(stream, "host", allowed_fields={"inst_name", "host_belong_app"})
    files = MemoryFiles()
    files.put("transfer/tmp/relation/source.xlsx", stream)
    context = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "import")
    task = TransferService.submit(
        owner=transfer_owner,
        kind="import",
        model_id="host",
        team_id=1,
        include_children=False,
        params={},
        authorization=context.snapshot,
        schema_hash=context.schema_hash,
        idempotency_key="relations",
        source_key="transfer/tmp/relation/source.xlsx",
        source_hash=info["sha256"],
    )
    TransferExecution.run(task.pk, files=files)
    task = TransferService.get(transfer_owner, task.pk)
    assert task.status == "partial_success"
    assert task.summary == {
        "created": 1,
        "updated": 0,
        "failed_rows": 0,
        "created_relations": 0 if existing_relation else 1,
        "failed_relations": 1,
        **({"existing_relations": 1} if existing_relation else {}),
    }
    edges.assert_called_once()
    assert edges.call_args.kwargs["dst_inst_uuid"] == "123e4567-e89b-42d3-a456-426614174001"


def test_all_invalid_rows_fail_with_downloadable_report_and_no_writes(transfer_owner, authorized, fake_graph, monkeypatch):
    from apps.cmdb.services.instance import InstanceManage
    from apps.cmdb.services.transfer_validation import inspect_workbook
    from apps.cmdb.views.transfer_task import task_data

    attrs = ATTRS + [{"attr_id": "count", "attr_name": "数量", "attr_type": "int"}]
    monkeypatch.setattr(ModelManage, "search_model_attr", lambda *a, **k: attrs)
    monkeypatch.setattr(ModelManage, "search_model_attr_v2", lambda *a, **k: attrs)
    monkeypatch.setattr(
        "apps.cmdb.utils.Import.build_unique_rule_context", lambda _: SimpleNamespace(unique_rules=[], attrs_by_id={"inst_name": ATTRS[0]})
    )
    fake_graph("apps.cmdb.services.transfer_import", query_entity=([], 0))
    writes = Mock()
    monkeypatch.setattr(InstanceManage, "instance_create", writes)
    book = openpyxl.Workbook()
    book.active.title = "host"
    for row in (["实例名", "数量"], ["str", "int"], ["inst_name", "count"], ["one", "PRIVATE-SENTINEL"]):
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    info = inspect_workbook(stream, "host", allowed_fields={"inst_name", "count"})
    files = MemoryFiles()
    files.put("transfer/tmp/invalid/source.xlsx", stream)
    context = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "import")
    task = TransferService.submit(
        owner=transfer_owner,
        kind="import",
        model_id="host",
        team_id=1,
        include_children=False,
        params={},
        authorization=context.snapshot,
        schema_hash=context.schema_hash,
        idempotency_key="invalid",
        source_key="transfer/tmp/invalid/source.xlsx",
        source_hash=info["sha256"],
    )
    TransferExecution.run(task.pk, files=files)
    task = TransferService.get(transfer_owner, task.pk)
    assert task.status == "failed"
    assert task.summary["failed_rows"] == 1
    assert task.processed_rows == 1
    assert "download_errors" in task_data(task)["available_actions"]
    writes.assert_not_called()
    report = openpyxl.load_workbook(io.BytesIO(files.objects[task.artifacts["errors"]["key"]])).active
    assert "PRIVATE-SENTINEL" not in str(list(report.values))


@pytest.mark.parametrize("denial", ["no_action", "no_teams", "no_rules", "hidden_model"])
def test_authorization_denials_never_load_asset_data(transfer_owner, authorized, monkeypatch, denial):
    from apps.cmdb.services.transfer_service import TransferError

    actor = {
        "username": transfer_owner.username,
        "domain": transfer_owner.domain,
        "is_superuser": False,
        "permission": {"cmdb": ["asset_info-View"]},
        "group_list": [{"id": 1}],
        "roles": [],
    }
    if denial == "no_action":
        actor["permission"] = {}
    monkeypatch.setattr("apps.cmdb.services.transfer_authorization.build_user_authorization_context", lambda _: actor)
    monkeypatch.setattr(
        "apps.cmdb.utils.permission_util.GroupUtils.get_user_authorized_child_groups", lambda **kw: [] if denial == "no_teams" else [1]
    )
    if denial == "no_rules":
        monkeypatch.setattr("apps.cmdb.utils.permission_util.get_permission_rules", lambda **kw: {})
    if denial == "hidden_model":
        monkeypatch.setattr("apps.cmdb.services.model_visibility.BusinessModelVisibility.resolve", lambda ids: {})
    attrs = Mock()
    monkeypatch.setattr(ModelManage, "search_model_attr", attrs)
    with pytest.raises(TransferError):
        TransferAuthorization.resolve(transfer_owner, 1, False, "host", "export")
    attrs.assert_not_called()


def test_free_tag_option_growth_does_not_invalidate_its_own_import(transfer_owner, authorized, monkeypatch):
    tag = {"attr_id": "tag", "attr_name": "标签", "attr_type": "tag", "option": {"mode": "free", "options": []}}
    monkeypatch.setattr(ModelManage, "search_model_attr", lambda *a, **k: ATTRS + [tag])
    before = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "import")
    tag["option"]["options"] = [{"key": "env", "value": "production"}]
    after = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "import")
    assert before.schema_hash == after.schema_hash
    tag["option"]["mode"] = "strict"
    strict = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "import")
    assert strict.schema_hash != after.schema_hash


def test_equivalent_permission_rule_order_keeps_task_authorization(transfer_owner, authorized, monkeypatch):
    rules = {"instance": [{"id": "one", "permission": ["View", "Operate"]}, {"id": "two", "permission": ["View"]}]}
    monkeypatch.setattr("apps.cmdb.utils.permission_util.get_permission_rules", lambda **kw: rules)
    first = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "export")
    rules["instance"].reverse()
    rules["instance"][1]["permission"].reverse()
    second = TransferAuthorization.resolve(transfer_owner, 1, False, "host", "export")
    assert first.snapshot == second.snapshot


def test_export_acceptance_does_not_wait_for_broker(transfer_owner, authorized, monkeypatch, django_capture_on_commit_callbacks):
    from threading import Event, Thread
    from time import monotonic

    from apps.cmdb.tasks.transfer import execute_transfer

    release = Event()
    entered = Event()
    threads = []

    def track_thread(**kwargs):
        thread = Thread(**kwargs)
        threads.append(thread)
        return thread

    monkeypatch.setattr("apps.cmdb.services.transfer_dispatch.Thread", track_thread)

    def delayed_publish(*args, **kwargs):
        entered.set()
        release.wait(0.4)

    # 模拟消息发布延迟；接纳请求不能继承 Broker 的等待时间。
    publish = Mock(side_effect=delayed_publish)
    monkeypatch.setattr(execute_transfer, "apply_async", publish)
    started = monotonic()
    try:
        with django_capture_on_commit_callbacks(execute=True):
            result = create_request(transfer_owner, {"model_id": "host", "attr_list": ["inst_name"], "scope": "all"})
        elapsed = monotonic() - started
        assert result.status_code == 202
        assert elapsed < 0.2, f"任务接纳被消息发布阻塞 {elapsed:.3f}s"
        assert entered.wait(1)
    finally:
        release.set()
        for thread in threads:
            thread.join(timeout=1)
            assert not thread.is_alive()
