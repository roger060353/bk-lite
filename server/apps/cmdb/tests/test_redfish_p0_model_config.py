"""Redfish P0 模型种子：存储控制器、电源，以及磁盘/网卡/整机补字段。"""

import os

import openpyxl
import pandas as pd
import pytest

pytestmark = pytest.mark.unit

XLSX = os.path.join(os.path.dirname(__file__), "..", "support-files", "model_config.xlsx")

STR_OPTION = '{"validation_type":"unrestricted","custom_regex":"","widget_type":"single_line"}'
INT_OPTION = '{"min_value": "", "max_value": ""}'

CONTROLLER_FIELDS = {
    "sc_id": "str",
    "sc_name": "str",
    "sc_vendor": "str",
    "sc_model": "str",
    "sc_sn": "str",
    "sc_firmware": "str",
    "health": "str",
    "self_device": "str",
}

PSU_FIELDS = {
    "psu_name": "str",
    "psu_vendor": "str",
    "psu_model": "str",
    "psu_sn": "str",
    "psu_capacity_watts": "int",
    "health": "str",
    "self_device": "str",
}


def _records(sheet):
    headers = [cell.value for cell in sheet[2]]
    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=3, values_only=True)
        if any(value is not None and str(value).strip() != "" for value in values)
    ]


def _attr_map(workbook, model_id):
    return {row["attr_id"]: row for row in _records(workbook[f"attr-{model_id}"])}


def test_redfish_p0_models_exist_under_hardware_components():
    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    models = {row["model_id"]: row for row in _records(workbook["models"])}

    assert models["storage_controller"]["classification_id"] == "hardware_components"
    assert models["psu"]["classification_id"] == "hardware_components"
    assert models["storage_controller"]["app_topo_layer"] == "infrastructure"
    assert models["psu"]["app_topo_layer"] == "infrastructure"
    assert models["storage_controller"]["model_name"] == "存储控制器"
    assert models["psu"]["model_name"] == "服务器电源"


def test_redfish_p0_fields_and_contains_edges():
    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    controller = _attr_map(workbook, "storage_controller")
    psu = _attr_map(workbook, "psu")
    disk = _attr_map(workbook, "disk")
    nic = _attr_map(workbook, "nic")
    server = _attr_map(workbook, "physcial_server")

    assert {key: controller[key]["attr_type"] for key in CONTROLLER_FIELDS} == CONTROLLER_FIELDS
    assert {key: psu[key]["attr_type"] for key in PSU_FIELDS} == PSU_FIELDS
    assert controller["sc_id"]["is_required"] is True
    assert psu["psu_name"]["is_required"] is True
    assert controller["self_device"]["key_attribute"] is True
    assert psu["self_device"]["key_attribute"] is True
    assert controller["sc_id"]["option"] == STR_OPTION
    assert psu["psu_capacity_watts"]["option"] == INT_OPTION
    assert "psu_input_watts" not in psu
    assert "fan" not in controller

    assert disk["health"]["attr_type"] == "str"
    assert disk["disk_life_percent"]["attr_type"] == "int"
    assert nic["nic_speed_mbps"]["attr_type"] == "int"
    assert nic["nic_iface"]["attr_id"] == "nic_iface"
    assert server["power_state"]["attr_type"] == "str"
    assert server["health"]["attr_type"] == "str"

    associations = {(row["src_model_id"], row["dst_model_id"], row["asst_id"], row["mapping"]) for row in _records(workbook["asso-physcial_server"])}
    assert ("physcial_server", "storage_controller", "contains", "1:n") in associations
    assert ("physcial_server", "psu", "contains", "1:n") in associations
    assert ("physcial_server", "disk", "contains", "1:n") in associations
    assert ("physcial_server", "nic", "contains", "1:n") in associations


def test_model_init_loader_sees_redfish_p0_sheets():
    sheets = pd.read_excel(XLSX, sheet_name=None, header=1)
    model_ids = set(sheets["models"]["model_id"].dropna().astype(str))
    assert {"storage_controller", "psu"} <= model_ids
    assert "attr-storage_controller" in sheets
    assert "attr-psu" in sheets
    assert "health" in set(sheets["attr-disk"]["attr_id"].dropna().astype(str))
    assert "nic_speed_mbps" in set(sheets["attr-nic"]["attr_id"].dropna().astype(str))
    assert "power_state" in set(sheets["attr-physcial_server"]["attr_id"].dropna().astype(str))
