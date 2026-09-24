import hashlib
import zipfile

import openpyxl
from rest_framework import serializers

from apps.cmdb.services.transfer_service import TransferError


class ExportRequest(serializers.Serializer):
    model_id = serializers.RegexField(r"^[A-Za-z_][A-Za-z0-9_]*$", max_length=128)
    scope = serializers.ChoiceField(choices=["all", "selected", "currentPage"])
    inst_uuids = serializers.ListField(child=serializers.UUIDField(format="hex_verbose"), max_length=100000, default=list)
    attr_list = serializers.ListField(child=serializers.CharField(max_length=128), max_length=256, allow_empty=False)
    association_list = serializers.ListField(child=serializers.CharField(max_length=256), max_length=128, default=list)

    def validate(self, data):
        if data["scope"] == "all" and data["inst_uuids"]:
            raise serializers.ValidationError("全部导出不能同时指定实例列表")
        if data["scope"] != "all" and not data["inst_uuids"]:
            raise serializers.ValidationError("请选择需要导出的实例")
        for key in ("inst_uuids", "attr_list", "association_list"):
            if len(data[key]) != len(set(data[key])):
                raise serializers.ValidationError("导出参数不能包含重复项")
        return data


def inspect_workbook(stream, model_id, *, allowed_fields):
    stream.seek(0, 2)
    if stream.tell() > 20 * 1024 * 1024:
        raise TransferError("file_too_large", "上传文件不能超过 20 MiB", 413)
    stream.seek(0)
    digest = hashlib.sha256()
    while chunk := stream.read(64 * 1024):
        digest.update(chunk)
    stream.seek(0)
    try:
        with zipfile.ZipFile(stream) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(member.file_size for member in members) > 200 * 1024 * 1024:
                raise TransferError("workbook_limit", "Excel 解压内容超过允许范围", 413)
            if any(member.flag_bits & 1 for member in members):
                raise TransferError("invalid_workbook", "不支持加密 Excel 文件")
        stream.seek(0)
        book = openpyxl.load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.worksheets[0]
            if sheet.title != model_id or (sheet.max_row or 0) > 10003 or (sheet.max_column or 0) > 256:
                raise TransferError("invalid_workbook", "模型不匹配，或 Excel 超过 1 万数据行 / 256 列")
            sheet.reset_dimensions()  # 不信任 ZIP 内声明的工作表范围，逐行计数真实内容。
            keys = []
            count = 0
            cells = 0
            for row_number, row in enumerate(sheet.iter_rows(), 1):
                cells += len(row)
                if row_number > 10003 or len(row) > 256 or cells > 1000000:
                    raise TransferError("workbook_limit", "Excel 行列或单元格总数超过上限", 413)
                if row_number == 3:
                    keys = [cell.value for cell in row]
                    if keys and keys[0] == "字段标识(请勿编辑)":
                        keys = keys[1:]
                    if not keys or any(key not in allowed_fields for key in keys) or len(set(keys)) != len(keys):
                        raise TransferError("invalid_workbook", "表头存在未知、重复或不可导入字段，请重新下载模板")
                if any(cell.data_type == "f" or len(str(cell.value or "")) > 32767 for cell in row):
                    raise TransferError("invalid_workbook", "不支持公式单元格或超长文本")
                if row_number >= 4 and any(cell.value is not None for cell in row):
                    count += 1
            if not keys or not count:
                raise TransferError("invalid_workbook", "文件没有可导入的数据")
        finally:
            book.close()
    except (zipfile.BadZipFile, KeyError, IndexError, ValueError, TypeError, OSError) as exc:
        raise TransferError("invalid_workbook", "文件不是有效的 Excel 模板") from exc
    finally:
        stream.seek(0)
    return {"rows": count, "sha256": digest.hexdigest()}
