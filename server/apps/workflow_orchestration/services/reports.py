from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate
from openpyxl import load_workbook
from xlsxjinja import BookWriter

MAX_TEMPLATE_BYTES = 5 * 1024 * 1024
MAX_TEMPLATE_ENTRIES = 200
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_TEMPLATE_TAGS = 200
MAX_LOOP_DEPTH = 10
CARBONE_MARKER = re.compile(r"\{d\.")
IDENT = re.compile(r"^[A-Za-z_][\w]*$")
ATTR_PATH = re.compile(r"^[A-Za-z_][\w]*(?:\.[A-Za-z_][\w]*)*$")
FOR_TAG = re.compile(
    r"^(?:(?:tr|tc|p|r)\s+)?for\s+([A-Za-z_][\w]*)\s+in\s+(.+?)(?:\s*if\s+.+)?$",
    re.IGNORECASE,
)
ENDFOR_TAG = re.compile(r"^(?:(?:tr|tc|p|r)\s+)?endfor\b", re.IGNORECASE)
DENIED_PATH_SEGMENTS = {"password", "passwd", "secret", "token", "credential", "private_key", "cookie", "authorization"}
_MISSING = object()


class ReportTemplateError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedReportTemplate:
    format: str
    placeholders: tuple[str, ...]
    loops: tuple[str, ...]
    required_metrics: tuple[str, ...]
    sha256: str
    size: int


def _validate_zip(content: bytes, fmt: str) -> None:
    if len(content) > MAX_TEMPLATE_BYTES:
        raise ReportTemplateError("模板超过 5 MiB 限额")
    try:
        archive = zipfile.ZipFile(BytesIO(content))
    except zipfile.BadZipFile as error:
        raise ReportTemplateError(f"{fmt} 不是有效的 Office Open XML 文件") from error
    infos = archive.infolist()
    if len(infos) > MAX_TEMPLATE_ENTRIES or sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_BYTES:
        raise ReportTemplateError("模板解压规模超过安全限额")
    for info in infos:
        name = info.filename.replace("\\", "/")
        if name.startswith("/") or ".." in name.split("/"):
            raise ReportTemplateError("模板包含非法文件路径")
        lowered = name.lower()
        if "vbaproject" in lowered or lowered.startswith("xl/externallinks/"):
            raise ReportTemplateError("模板不得包含宏或外部数据连接")
        if info.compress_size and info.file_size / info.compress_size > 100:
            raise ReportTemplateError("模板包含异常压缩条目")
    for info in infos:
        if info.filename.endswith(".rels"):
            payload = archive.read(info).decode("utf-8", errors="ignore").lower()
            if 'targetmode="external"' in payload:
                raise ReportTemplateError("模板不得包含外部链接")


def _document_text(content: bytes, fmt: str) -> str:
    if fmt == "docx":
        document = Document(BytesIO(content))
        chunks: list[str] = []
        for child in document.element.body:
            tag = child.tag.split("}")[-1]
            if tag == "p":
                chunks.append(Paragraph(child, document).text)
            elif tag == "tbl":
                table = Table(child, document)
                chunks.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(chunks)
    if fmt == "xlsx":
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
        return "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str))
    raise ReportTemplateError("只支持 docx 或 xlsx 模板")


def _expr_path(expr: str) -> str | None:
    cleaned = expr.split("|", 1)[0].strip()
    if not cleaned or not ATTR_PATH.match(cleaned):
        return None
    return cleaned


def _path_segments(path: str) -> list[str]:
    return [part for part in path.split(".") if part]


def _deny_sensitive(path: str) -> None:
    denied = sorted({segment.lower() for segment in _path_segments(path)}.intersection(DENIED_PATH_SEGMENTS))
    if denied:
        raise ReportTemplateError(f"模板不得引用敏感字段: {path}")


def _resolve_iterable(raw: str, bindings: dict[str, str]) -> str:
    path = _expr_path(raw)
    if path is None:
        raise ReportTemplateError(f"模板循环语法错误: {raw}")
    parts = _path_segments(path)
    if not parts:
        raise ReportTemplateError(f"模板循环语法错误: {raw}")
    root = parts[0]
    if root in bindings:
        return ".".join([bindings[root], *parts[1:]])
    return path


def _absolute_placeholder(path: str, bindings: dict[str, str]) -> str:
    parts = _path_segments(path)
    if not parts:
        return path
    root = parts[0]
    if root in bindings:
        return ".".join([bindings[root], *parts[1:]]) if len(parts) > 1 else bindings[root]
    return path


def parse_report_template(content: bytes, fmt: str) -> ParsedReportTemplate:
    normalized = fmt.lower().lstrip(".")
    if normalized not in {"docx", "xlsx"}:
        raise ReportTemplateError("只支持 docx 或 xlsx 模板")
    _validate_zip(content, normalized)
    text = _document_text(content, normalized)
    if CARBONE_MARKER.search(text):
        raise ReportTemplateError("模板请使用 Jinja 语法 {{ 字段 }} / {% for %}，不支持 Carbone {d.字段}")

    tokens = re.findall(r"(\{\{.*?\}\}|\{%.*?%\})", text, flags=re.DOTALL)
    if len(tokens) > MAX_TEMPLATE_TAGS:
        raise ReportTemplateError("模板占位符数量超过 200 个")

    bindings: dict[str, str] = {}
    loop_stack: list[str] = []
    loops: list[str] = []
    labels: list[str] = []

    for token in tokens:
        if token.startswith("{{"):
            path = _expr_path(token[2:-2].strip())
            if path is None:
                continue
            absolute = _absolute_placeholder(path, bindings)
            _deny_sensitive(absolute)
            labels.append(absolute)
            continue

        tag = " ".join(token[2:-2].split())
        for_match = FOR_TAG.match(tag)
        if for_match:
            loop_var = for_match.group(1)
            if not IDENT.match(loop_var):
                raise ReportTemplateError(f"模板循环变量非法: {loop_var}")
            iterable = _resolve_iterable(for_match.group(2), bindings)
            _deny_sensitive(iterable)
            if len(loop_stack) >= MAX_LOOP_DEPTH:
                raise ReportTemplateError("模板循环嵌套超过 10 层")
            if iterable not in loops:
                loops.append(iterable)
            loop_stack.append(loop_var)
            bindings[loop_var] = iterable
            continue
        if ENDFOR_TAG.match(tag):
            if not loop_stack:
                raise ReportTemplateError("模板循环缺少开始标记")
            finished = loop_stack.pop()
            bindings.pop(finished, None)

    if loop_stack:
        raise ReportTemplateError(f"模板循环缺少结束标记: {bindings.get(loop_stack[-1], loop_stack[-1])}")

    if len(loops) > 20:
        raise ReportTemplateError("模板循环数量超过 20 个")
    placeholders = tuple(dict.fromkeys(labels))
    return ParsedReportTemplate(
        format=normalized,
        placeholders=placeholders,
        loops=tuple(loops),
        required_metrics=placeholders,
        sha256=hashlib.sha256(content).hexdigest(),
        size=len(content),
    )


def _lookup(data: Any, parts: list[str]) -> Any:
    value = data
    for name in parts:
        if isinstance(value, list):
            value = value[0] if value else _MISSING
        if value is _MISSING:
            return _MISSING
        if not isinstance(value, dict) or name not in value:
            return _MISSING
        value = value[name]
    return value


def missing_template_fields(parsed: ParsedReportTemplate, data: dict[str, Any]) -> list[str]:
    missing: list[str] = []
    for label in parsed.placeholders:
        parts = _path_segments(label)
        if not parts:
            continue
        if _lookup(data, parts) is _MISSING:
            missing.append(label)
    return missing


def _render_docx(content: bytes, data: dict[str, Any]) -> bytes:
    template = DocxTemplate(BytesIO(content))
    try:
        template.render(data)
    except Exception as error:  # noqa: BLE001 - surface engine failures as template errors
        raise ReportTemplateError(f"Word 模板渲染失败: {error}") from error
    output = BytesIO()
    template.save(output)
    rendered = output.getvalue()
    if not rendered:
        raise ReportTemplateError("Word 模板渲染返回了空文档")
    return rendered


def _render_xlsx(content: bytes, data: dict[str, Any]) -> bytes:
    try:
        writer = BookWriter(BytesIO(content))
        sheet_names = [state.name for state in writer.sheet_resource_map.sheet_state_list]
        payloads = [
            {
                **data,
                "tpl_name": sheet_name,
                "sheet_name": sheet_name,
            }
            for sheet_name in sheet_names
        ] or [data]
        writer.render_book(payloads)
        output = BytesIO()
        writer.save(output)
    except Exception as error:  # noqa: BLE001 - surface engine failures as template errors
        raise ReportTemplateError(f"Excel 模板渲染失败: {error}") from error
    rendered = output.getvalue()
    if not rendered:
        raise ReportTemplateError("Excel 模板渲染返回了空文档")
    return rendered


def render_report(content: bytes, fmt: str, data: dict[str, Any]) -> bytes:
    parsed = parse_report_template(content, fmt)
    if not isinstance(data, dict):
        raise ReportTemplateError("文档数据必须是 JSON 对象")
    if parsed.format == "docx":
        return _render_docx(content, data)
    return _render_xlsx(content, data)
