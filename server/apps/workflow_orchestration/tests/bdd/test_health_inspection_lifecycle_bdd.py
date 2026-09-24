from __future__ import annotations

from io import BytesIO, StringIO
from pathlib import Path

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from docx import Document
from openpyxl import load_workbook
from pytest_bdd import given, scenarios, then, when

from apps.workflow_orchestration.models import Workflow, WorkflowExecution, WorkflowInteraction
from apps.workflow_orchestration.services.reports import parse_report_template, render_report

FEATURE = str(Path(__file__).parent / "health_inspection_lifecycle.feature")
TEMPLATE_DIR = Path(__file__).resolve().parents[5] / "web/public/workflow-orchestration/templates"
scenarios(FEATURE)

pytestmark = [pytest.mark.bdd, pytest.mark.django_db]


@pytest.fixture
def ctx(authenticated_user, mocker):
    conductor = mocker.patch("apps.workflow_orchestration.management.commands.seed_workflow_orchestration_demo.ConductorClient").return_value
    mocker.patch(
        "apps.workflow_orchestration.management.commands.seed_workflow_orchestration_demo.seed_builtin_health_template_snapshot",
        side_effect=lambda fmt, team_id, store=None: {
            "object_key": f"workflow-orchestration/templates/demo/team-{team_id}/health.{fmt}",
            "format": fmt,
            "sha256": "a" * 64,
            "size": 128,
            "filename_prefix": f"health-inspection-{fmt}",
        },
    )
    return {"user": authenticated_user, "team_id": 1, "error": None, "conductor": conductor}


def _run_seed(ctx, *, confirm=True, username=None):
    try:
        call_command(
            "seed_workflow_orchestration_demo",
            team_id=ctx["team_id"],
            username=username or ctx["user"].username,
            domain=ctx["user"].domain,
            confirm=confirm,
            stdout=StringIO(),
        )
    except CommandError as error:
        ctx["error"] = error


def _rendered_text(fmt: str, content: bytes) -> str:
    if fmt == "docx":
        document = Document(BytesIO(content))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        chunks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(chunks)
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    return "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)


@given("页面查看用户有权访问当前团队")
def authorized_viewer(ctx):
    ctx["user"].group_list = [{"id": ctx["team_id"], "name": "Default"}]
    ctx["user"].save(update_fields=("group_list",))


@given("页面查看用户只有其他团队权限")
def unauthorized_viewer(ctx):
    ctx["user"].group_list = [{"id": 2, "name": "Other"}]
    ctx["user"].save(update_fields=("group_list",))


@given("当前团队已有一条非演示流程")
def business_workflow(ctx):
    ctx["business_workflow"] = Workflow.objects.create(name="用户业务流程", team=[ctx["team_id"]])


@when("管理员确认重建 MVP 验收数据")
def confirmed_seed(ctx):
    _run_seed(ctx)


@when("管理员未确认就重建 MVP 验收数据")
def unconfirmed_seed(ctx):
    _run_seed(ctx, confirm=False)


@when("不存在的页面查看用户尝试重建 MVP 验收数据")
def unknown_viewer_seed(ctx):
    _run_seed(ctx, username="missing-user")


@when("管理员连续两次确认重建 MVP 验收数据")
def seed_twice(ctx):
    _run_seed(ctx)
    _run_seed(ctx)


@then("应当生成七条已发布流程和十一条执行记录")
def seven_workflows_and_eleven_executions(ctx):
    assert ctx["error"] is None
    assert Workflow.objects.filter(name__startswith="[TDD/BDD]", status=Workflow.Status.PUBLISHED).count() == 7
    assert WorkflowExecution.objects.filter(workflow__name__startswith="[TDD/BDD]").count() == 11


@then("应当有一条待审批记录")
def one_pending_approval():
    assert WorkflowInteraction.objects.filter(status=WorkflowInteraction.Status.PENDING).count() == 1


@then("应当同时存在 Word 与 Excel 健康巡检流程")
def word_and_excel_health_workflows():
    names = set(Workflow.objects.filter(name__startswith="[TDD/BDD]").values_list("name", flat=True))
    assert "[TDD/BDD] 主机健康巡检 Word" in names
    assert "[TDD/BDD] 主机健康巡检 Excel" in names


@then("健康巡检数据应当包含 CPU 内存 磁盘和网络维度")
def multidimensional_health_data(ctx):
    execution = WorkflowExecution.objects.get(trigger_id="health-word-success")
    ctx["health_data"] = execution.output["job"]
    categories = {metric["category"] for metric in ctx["health_data"]["results"][0]["data"]["metrics"]}
    assert categories >= {"CPU", "内存", "磁盘", "网络"}


@then("Excel 和 Word 内置模板都应当能渲染健康巡检数据")
def builtin_templates_render(ctx):
    for fmt in ("xlsx", "docx"):
        content = (TEMPLATE_DIR / f"health-inspection-example.{fmt}").read_bytes()
        parsed = parse_report_template(content, fmt)
        rendered = render_report(content, fmt, ctx["health_data"])
        text = _rendered_text(fmt, rendered)
        assert "results" in parsed.loops
        assert "10.10.90.120" in text
        assert "CPU Total" in text
        assert "Ethernet0" in text
        assert "{{" not in text


@then("命令应当拒绝执行并且不生成流程")
def unconfirmed_rejected(ctx):
    assert "显式传入 --confirm" in str(ctx["error"])
    assert not Workflow.objects.exists()


@then("命令应当拒绝未知用户并且不生成流程")
def unknown_viewer_rejected(ctx):
    assert "找不到页面查看用户" in str(ctx["error"])
    assert not Workflow.objects.exists()


@then("命令应当拒绝跨团队操作并且不生成流程")
def cross_team_rejected(ctx):
    assert "无权访问团队" in str(ctx["error"])
    assert not Workflow.objects.exists()


@then("非演示流程应当保留且演示数据不应当重复")
def scoped_idempotent_rebuild(ctx):
    assert Workflow.objects.filter(pk=ctx["business_workflow"].pk, name="用户业务流程").exists()
    assert Workflow.objects.filter(name__startswith="[TDD/BDD]").count() == 7
    assert WorkflowExecution.objects.filter(workflow__name__startswith="[TDD/BDD]").count() == 11


@then("多磁盘告警执行应当标记告警并保留具体磁盘维度")
def warning_evidence():
    execution = WorkflowExecution.objects.get(trigger_id="health-warning")
    metrics = execution.output["job"]["results"][0]["data"]["metrics"]
    disk_warning = next(metric for metric in metrics if metric["category"] == "磁盘" and metric["object_name"] == "D:")
    assert execution.has_warnings is True
    assert execution.warning_count == 1
    assert disk_warning["dimensions"] == "mount=D:"
    assert disk_warning["health_status"] == "WARNING"


@then("Word 与 Excel 巡检流程应当分别挂载 docx 与 xlsx 模板快照")
def health_template_snapshots():
    word = Workflow.objects.get(name="[TDD/BDD] 主机健康巡检 Word")
    excel = Workflow.objects.get(name="[TDD/BDD] 主机健康巡检 Excel")
    word_report = next(task for task in word.definition["tasks"] if task["name"] == "bklite_document_render")
    excel_report = next(task for task in excel.definition["tasks"] if task["name"] == "bklite_document_render")
    assert word_report["inputParameters"]["template_snapshot"]["format"] == "docx"
    assert excel_report["inputParameters"]["template_snapshot"]["format"] == "xlsx"
