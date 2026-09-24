from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from apps.workflow_orchestration.services.reports import ReportTemplateError, missing_template_fields, parse_report_template, render_report


def _docx_bytes():
    document = Document()
    document.add_paragraph("主机总数：{{ summary.total }}")
    table = document.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "主机"
    table.cell(0, 1).text = "状态"
    table.cell(1, 0).text = "{%tr for r in results %}"
    table.cell(2, 0).text = "{{ r.target.name }}"
    table.cell(2, 1).text = "{{ r.status }}"
    table.cell(3, 0).text = "{%tr endfor %}"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "主机总数"
    sheet["B1"] = "{{ summary.total }}"
    sheet["A2"] = "{% for r in results %}"
    sheet["A3"] = "{{ r.target.name }}"
    sheet["B3"] = "{{ r.target.ip }}"
    sheet["A4"] = "{% endfor %}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.parametrize("fmt,builder", [("docx", _docx_bytes), ("xlsx", _xlsx_bytes)])
def test_template_parse_and_renders_through_jinja_engines(fmt, builder):
    content = builder()
    parsed = parse_report_template(content, fmt)
    data = {"summary": {"total": 1}, "results": [{"target": {"name": "server-01", "ip": "10.0.0.1"}, "status": "SUCCESS"}]}

    rendered = render_report(content, fmt, data)

    assert "summary.total" in parsed.placeholders
    assert "results" in parsed.loops
    assert missing_template_fields(parsed, data) == []
    if fmt == "docx":
        text = "\n".join(
            [paragraph.text for paragraph in Document(BytesIO(rendered)).paragraphs]
            + [cell.text for table in Document(BytesIO(rendered)).tables for row in table.rows for cell in row.cells]
        )
    else:
        workbook = load_workbook(BytesIO(rendered))
        text = "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert "server-01" in text
    assert "{{" not in text


def test_unknown_sensitive_field_is_rejected():
    workbook = Workbook()
    workbook.active["A1"] = "{% for r in results %}"
    workbook.active["A2"] = "{{ r.password }}"
    workbook.active["A3"] = "{% endfor %}"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ReportTemplateError, match="password"):
        parse_report_template(output.getvalue(), "xlsx")


def test_legacy_carbone_syntax_is_rejected():
    workbook = Workbook()
    workbook.active["A1"] = "{d.summary.total}"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ReportTemplateError, match="Jinja"):
        parse_report_template(output.getvalue(), "xlsx")


def test_unclosed_jinja_loop_is_rejected():
    workbook = Workbook()
    workbook.active["A1"] = "{% for r in results %}"
    workbook.active["A2"] = "{{ r.target.ip }}"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ReportTemplateError, match="循环"):
        parse_report_template(output.getvalue(), "xlsx")
