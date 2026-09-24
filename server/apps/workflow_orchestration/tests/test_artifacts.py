import hashlib
from datetime import timedelta
from io import BytesIO
from pathlib import Path

import pytest
from django.utils import timezone
from docx import Document
from openpyxl import Workbook

from apps.workflow_orchestration.models import ExecutionArtifact, Workflow, WorkflowExecution
from apps.workflow_orchestration.serializers import ExecutionArtifactSerializer
from apps.workflow_orchestration.services.artifacts import cleanup_expired_artifacts
from apps.workflow_orchestration.services.atoms import report_atom


class MemoryStore:
    def __init__(self):
        self.objects = {}

    def put(self, key, content):
        self.objects[key] = content.read()

    def get(self, key):
        content = self.objects[key]
        return content, key.rsplit("/", 1)[-1], len(content)


def _inspection():
    return {
        "summary": {
            "total_hosts": 1,
            "successful_hosts": 1,
            "failed_hosts": 0,
            "abnormal_hosts": 1,
            "execution_outcome": "completed_with_warnings",
        },
        "hosts": [
            {
                "node_id": "node-1",
                "host_name": "server-01",
                "ip": "10.0.0.1",
                "os_type": "linux",
                "os_version": "Ubuntu",
                "architecture": "x86_64",
                "collection_status": "success",
                "health_status": "warning",
                "cpu": {"cores": 4, "usage_percent": 10},
                "memory": {"total_bytes": 100, "used_bytes": 20, "available_bytes": 80, "usage_percent": 20},
                "disks": [
                    {
                        "name": "/dev/sda",
                        "mount_point": "/",
                        "total_bytes": 100,
                        "used_bytes": 40,
                        "available_bytes": 60,
                        "usage_percent": 40,
                        "health": {"status": "healthy"},
                    },
                    {
                        "name": "/dev/sdb",
                        "mount_point": "/data",
                        "total_bytes": 100,
                        "used_bytes": 85,
                        "available_bytes": 15,
                        "usage_percent": 85,
                        "health": {"status": "warning"},
                    },
                ],
            }
        ],
    }


def _template(fmt):
    output = BytesIO()
    if fmt == "docx":
        document = Document()
        table = document.add_table(rows=6, cols=2)
        table.cell(0, 0).text = "{%tr for r in results %}"
        table.cell(1, 0).text = "{{ r.target.name }}"
        table.cell(1, 1).text = "{{ r.target.ip }}"
        table.cell(2, 0).text = "{%tr for m in r.data.metrics %}"
        table.cell(3, 0).text = "{{ m.metric_name }}"
        table.cell(3, 1).text = "{{ m.value }}"
        table.cell(4, 0).text = "{%tr endfor %}"
        table.cell(5, 0).text = "{%tr endfor %}"
        document.save(output)
    else:
        workbook = Workbook()
        sheet = workbook.active
        for row in (
            ["{% for r in results %}"],
            ["{{ r.target.name }}", "{{ r.target.ip }}"],
            ["{% for m in r.data.metrics %}"],
            ["{{ m.metric_name }}", "{{ m.value }}"],
            ["{% endfor %}"],
            ["{% endfor %}"],
        ):
            sheet.append(row)
        workbook.save(output)
    return output.getvalue()


@pytest.mark.django_db
def test_cleanup_deletes_object_but_keeps_audit_summary(mocker):
    workflow = Workflow.objects.create(name="inspection", team=[7], definition={})
    execution = WorkflowExecution.objects.create(workflow=workflow, workflow_version=1, team=[7])
    artifact = ExecutionArtifact.objects.create(
        execution=execution,
        team=[7],
        format="docx",
        object_key="reports/a.docx",
        filename="a.docx",
        content_type="application/docx",
        sha256="a" * 64,
        size=10,
        summary={"total_hosts": 2},
        expires_at=timezone.now() - timedelta(seconds=1),
    )
    store = mocker.Mock()

    result = cleanup_expired_artifacts(store=store)

    artifact.refresh_from_db()
    store.delete.assert_called_once_with("reports/a.docx")
    assert result == {"selected": 1, "deleted": 1, "failed": 0}
    assert artifact.deleted_at is not None
    assert artifact.summary == {"total_hosts": 2}


@pytest.mark.django_db
def test_cleanup_leaves_artifact_reclaimable_when_object_delete_fails(mocker):
    workflow = Workflow.objects.create(name="inspection-fail", team=[7], definition={})
    execution = WorkflowExecution.objects.create(workflow=workflow, workflow_version=1, team=[7])
    artifact = ExecutionArtifact.objects.create(
        execution=execution,
        team=[7],
        format="docx",
        object_key="reports/fail.docx",
        filename="fail.docx",
        content_type="application/docx",
        sha256="b" * 64,
        size=10,
        summary={},
        expires_at=timezone.now() - timedelta(seconds=1),
    )
    store = mocker.Mock()
    store.delete.side_effect = RuntimeError("s3 down")

    result = cleanup_expired_artifacts(store=store)

    artifact.refresh_from_db()
    assert result == {"selected": 1, "deleted": 0, "failed": 1}
    assert artifact.deleted_at is None


@pytest.mark.django_db
def test_report_atom_uses_worker_trusted_context_without_ui_execution_bindings():
    """UI 发布的文档节点只有 data/template_snapshot；执行身份由 Worker 注入。"""
    workflow = Workflow.objects.create(name="inspection-ui", team=[7], definition={})
    execution = WorkflowExecution.objects.create(workflow=workflow, workflow_version=1, team=[7])
    store = MemoryStore()
    template = _template("docx")
    template_key = "workflow-orchestration/templates/frozen/ui-template.docx"
    store.objects[template_key] = template

    output = report_atom(
        {
            "template_snapshot": {
                "object_key": template_key,
                "format": "docx",
                "sha256": hashlib.sha256(template).hexdigest(),
                "size": len(template),
                "filename_prefix": "inspection",
            },
            "data": _inspection(),
            "__bklite_context": {
                "execution_id": str(execution.id),
                "organization_id": 7,
                "actor": {"username": "alice", "domain": "example.com"},
            },
        },
        store=store,
    )

    artifact = ExecutionArtifact.objects.get(pk=output["artifact"]["id"])
    assert artifact.execution_id == execution.id
    assert artifact.team == [7]
    assert artifact.format == "docx"
    assert store.objects[artifact.object_key]


@pytest.mark.django_db
@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_report_atom_persists_downloadable_office_artifact(fmt):
    workflow = Workflow.objects.create(name="inspection", team=[7], definition={})
    execution = WorkflowExecution.objects.create(workflow=workflow, workflow_version=1, team=[7])
    store = MemoryStore()
    template = _template(fmt)
    template_key = f"workflow-orchestration/templates/frozen/template.{fmt}"
    store.objects[template_key] = template

    output = report_atom(
        {
            "execution_id": str(execution.id),
            "team": 7,
            "template_snapshot": {
                "object_key": template_key,
                "format": fmt,
                "sha256": hashlib.sha256(template).hexdigest(),
                "size": len(template),
                "filename_prefix": "inspection",
            },
            "data": _inspection(),
        },
        store=store,
    )

    artifact = ExecutionArtifact.objects.get(pk=output["artifact"]["id"])
    content = store.objects[artifact.object_key]
    assert artifact.sha256
    assert artifact.size == len(content)
    assert artifact.summary["execution_outcome"] == "completed_with_warnings"
    assert artifact.expires_at > timezone.now() + timedelta(days=29)
    assert content
    assert content != template
    assert isinstance(output["field_warnings"], list)


@pytest.mark.django_db
def test_report_atom_merges_additional_job_envelope_for_multi_host_report():
    workflow = Workflow.objects.create(name="inspection-merge", team=[7], definition={})
    execution = WorkflowExecution.objects.create(workflow=workflow, workflow_version=1, team=[7])
    store = MemoryStore()
    template = _template("xlsx")
    template_key = "workflow-orchestration/templates/frozen/template.xlsx"
    store.objects[template_key] = template

    primary = {
        "results": [
            {
                "target": {"name": "win", "ip": "10.10.90.120", "operating_system": "windows"},
                "status": "SUCCESS",
                "data": {
                    "conclusion": "健康",
                    "metrics": [{"metric_name": "cpu", "value": 1}],
                    "critical": [],
                    "warning": [],
                    "normal": [],
                },
            }
        ],
        "summary": {"total": 1, "succeeded": 1, "failed": 0},
        "job_task_ids": [101],
    }
    additional = {
        "results": [
            {
                "target": {"name": "linux", "ip": "192.168.64.5", "operating_system": "linux"},
                "status": "SUCCESS",
                "data": {
                    "conclusion": "健康",
                    "metrics": [{"metric_name": "mem", "value": 2}],
                    "critical": [],
                    "warning": [],
                    "normal": [],
                },
            }
        ],
        "summary": {"total": 1, "succeeded": 1, "failed": 0},
        "job_task_ids": [102],
    }

    output = report_atom(
        {
            "execution_id": str(execution.id),
            "team": 7,
            "template_snapshot": {
                "object_key": template_key,
                "format": "xlsx",
                "sha256": hashlib.sha256(template).hexdigest(),
                "size": len(template),
                "filename_prefix": "inspection",
            },
            "data": primary,
            "additional_data": additional,
        },
        store=store,
    )

    artifact = ExecutionArtifact.objects.get(pk=output["artifact"]["id"])
    assert artifact.summary == {"total": 2, "succeeded": 2, "failed": 0}
    content = store.objects[artifact.object_key]
    assert artifact.size == len(content)
    assert content != template


@pytest.mark.django_db
def test_report_atom_uses_published_template_snapshot():
    workflow = Workflow.objects.create(name="inspection", team=[7], definition={})
    execution = WorkflowExecution.objects.create(workflow=workflow, workflow_version=3, team=[7])
    store = MemoryStore()
    frozen = _template("docx")
    store.objects["workflow-orchestration/templates/frozen/template.docx"] = frozen

    output = report_atom(
        {
            "execution_id": str(execution.id),
            "team": 7,
            "template_snapshot": {
                "object_key": "workflow-orchestration/templates/frozen/template.docx",
                "format": "docx",
                "sha256": hashlib.sha256(frozen).hexdigest(),
                "size": len(frozen),
            },
            "data": _inspection(),
        },
        store=store,
    )

    assert output["artifact"]["format"] == "docx"


@pytest.mark.django_db
def test_artifact_download_url_stays_on_authenticated_api_client_path():
    workflow = Workflow.objects.create(name="inspection", team=[7], definition={})
    execution = WorkflowExecution.objects.create(workflow=workflow, workflow_version=1, team=[7])
    artifact = ExecutionArtifact.objects.create(
        execution=execution,
        team=[7],
        format="docx",
        object_key="reports/a.docx",
        filename="a.docx",
        content_type="application/docx",
        sha256="a" * 64,
        size=10,
        expires_at=timezone.now() + timedelta(days=1),
    )

    assert ExecutionArtifactSerializer(artifact).data["download_url"] == (f"/workflow_orchestration/api/artifacts/{artifact.id}/download/")


def test_production_supervisor_schedules_artifact_cleanup():
    config = Path("support-files/release/supervisor/workflow_orchestration_worker.conf").read_text()

    assert "[program:workflow_orchestration_artifact_cleanup]" in config
    assert "run_workflow_artifact_cleanup_scheduler" in config
