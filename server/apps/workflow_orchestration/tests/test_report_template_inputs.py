from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from apps.workflow_orchestration.models import Workflow, WorkflowExecution
from apps.workflow_orchestration.services.report_template_inputs import (
    resolve_uploaded_report_template,
    store_uploaded_report_template,
    verify_template_content,
)


class MemoryStore:
    def __init__(self):
        self.objects = {}

    def put(self, key, content):
        self.objects[key] = content.read()


def _docx_template():
    document = Document()
    document.add_paragraph("结果：{{ summary.total }}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.django_db
def test_uploaded_template_receipt_is_bound_to_workflow_version_team_and_actor():
    workflow = Workflow.objects.create(name="inspection", team=[7], definition={})
    execution = WorkflowExecution.objects.create(
        workflow=workflow,
        workflow_version=3,
        team=[7],
        started_by="alice",
        domain="example.com",
    )
    store = MemoryStore()
    content = _docx_template()

    reference = store_uploaded_report_template(
        SimpleUploadedFile("custom.docx", content),
        workflow_id=workflow.id,
        workflow_version=3,
        team=7,
        username="alice",
        domain="example.com",
        store=store,
    )
    snapshot = resolve_uploaded_report_template(reference, execution)

    verify_template_content(store.objects[snapshot["object_key"]], snapshot)
    assert reference["kind"] == "uploaded"
    assert reference["name"] == "custom.docx"
    assert len(reference["token"]) <= 8192
    assert "placeholders" not in snapshot
    assert snapshot["object_key"].startswith("workflow-orchestration/templates/uploads/7/")

    execution.started_by = "mallory"
    with pytest.raises(ValueError, match="不匹配"):
        resolve_uploaded_report_template(reference, execution)


def test_uploaded_template_rejects_non_office_content():
    with pytest.raises(ValueError):
        store_uploaded_report_template(
            SimpleUploadedFile("payload.txt", b"not an office template"),
            workflow_id=1,
            workflow_version=1,
            team=7,
            username="alice",
            domain="example.com",
            store=MemoryStore(),
        )


def test_uploaded_template_applies_field_format_and_size_constraints():
    content = _docx_template()
    with pytest.raises(ValueError, match="文件类型不受支持"):
        store_uploaded_report_template(
            SimpleUploadedFile("custom.docx", content),
            workflow_id=1,
            workflow_version=1,
            team=7,
            username="alice",
            domain="example.com",
            allowed_formats=("xlsx",),
            store=MemoryStore(),
        )
    with pytest.raises(ValueError, match="文件超过"):
        store_uploaded_report_template(
            SimpleUploadedFile("custom.docx", content),
            workflow_id=1,
            workflow_version=1,
            team=7,
            username="alice",
            domain="example.com",
            max_bytes=10,
            store=MemoryStore(),
        )


def test_publish_requires_a_document_template_snapshot():
    from apps.workflow_orchestration.services.definitions import build_health_inspection_definition
    from apps.workflow_orchestration.services.report_template_inputs import freeze_document_templates

    with pytest.raises(ValueError, match="必须上传"):
        freeze_document_templates(build_health_inspection_definition())
