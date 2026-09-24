from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from apps.workflow_orchestration.services.reports import parse_report_template, render_report

TEMPLATE_DIR = Path(__file__).resolve().parents[4] / "web/public/workflow-orchestration/templates"


def _metric(
    *,
    category: str,
    object_name: str,
    metric_name: str,
    value: float,
    unit: str,
    health_status: str,
    detail: str,
    collected_at: str,
) -> dict:
    return {
        "category": category,
        "object_type": "resource",
        "object_name": object_name,
        "dimensions": f"name={object_name}",
        "metric_name": metric_name,
        "value": value,
        "unit": unit,
        "warning_threshold": 80,
        "critical_threshold": 90,
        "health_status": health_status,
        "collected_at": collected_at,
        "detail": detail,
    }


def _host_result(
    *,
    name: str,
    ip: str,
    operating_system: str,
    conclusion: str,
    metrics: list[dict],
    collected_at: str,
) -> dict:
    critical = [item for item in metrics if item["health_status"] == "CRITICAL"]
    warning = [item for item in metrics if item["health_status"] == "WARNING"]
    normal = [item for item in metrics if item["health_status"] == "NORMAL"]
    return {
        "target": {
            "name": name,
            "ip": ip,
            "operating_system": operating_system,
        },
        "status": "SUCCESS",
        "data": {
            "collected_at": collected_at,
            "conclusion": conclusion,
            "metric_count": len(metrics),
            "critical_count": len(critical),
            "warning_count": len(warning),
            "normal_count": len(normal),
            "critical": critical,
            "warning": warning,
            "normal": normal,
            "metrics": metrics,
        },
        "error": None,
    }


HEALTH_REPORT_DATA = {
    "summary": {"total": 1, "succeeded": 1, "failed": 0},
    "results": [
        _host_result(
            name="job-web3",
            ip="10.10.90.120",
            operating_system="windows",
            conclusion="健康",
            collected_at="2026-09-18 10:30:00",
            metrics=[
                _metric(
                    category="CPU",
                    object_name="CPU Total",
                    metric_name="usage_percent",
                    value=42.3,
                    unit="%",
                    health_status="NORMAL",
                    detail="CPU 使用率正常",
                    collected_at="2026-09-18 10:30:00",
                ),
                _metric(
                    category="磁盘",
                    object_name="C:",
                    metric_name="usage_percent",
                    value=73.1,
                    unit="%",
                    health_status="NORMAL",
                    detail="系统盘剩余空间充足",
                    collected_at="2026-09-18 10:30:00",
                ),
                _metric(
                    category="网络",
                    object_name="Ethernet0",
                    metric_name="throughput_mbps",
                    value=18.6,
                    unit="Mbps",
                    health_status="NORMAL",
                    detail="网卡接收速率正常",
                    collected_at="2026-09-18 10:30:00",
                ),
            ],
        )
    ],
}


MULTI_HOST_REPORT_DATA = {
    "summary": {"total": 2, "succeeded": 2, "failed": 0},
    "results": [
        _host_result(
            name="job-web3",
            ip="10.10.90.120",
            operating_system="windows",
            conclusion="健康",
            collected_at="2026-09-23 18:00:00",
            metrics=[
                _metric(
                    category="CPU",
                    object_name="CPU Total",
                    metric_name="usage_percent",
                    value=12.0,
                    unit="%",
                    health_status="NORMAL",
                    detail="Windows CPU",
                    collected_at="2026-09-23 18:00:00",
                )
            ],
        ),
        _host_result(
            name="job-lab-linux-yum",
            ip="192.168.64.5",
            operating_system="linux",
            conclusion="需关注",
            collected_at="2026-09-23 18:00:01",
            metrics=[
                _metric(
                    category="磁盘",
                    object_name="/mnt/lima-cidata",
                    metric_name="usage_percent",
                    value=100.0,
                    unit="%",
                    health_status="CRITICAL",
                    detail="Linux 磁盘已满",
                    collected_at="2026-09-23 18:00:01",
                )
            ],
        ),
    ],
}


def _rendered_text(fmt: str, content: bytes) -> str:
    if fmt == "docx":
        document = Document(BytesIO(content))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        chunks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(chunks)
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    return "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_builtin_health_report_template_matches_job_output_contract_and_renders_multidimensional_metrics(fmt):
    template_path = TEMPLATE_DIR / f"health-inspection-example.{fmt}"
    content = template_path.read_bytes()

    parsed = parse_report_template(content, fmt)
    rendered = render_report(content, fmt, HEALTH_REPORT_DATA)

    assert "results" in parsed.loops
    assert "results.data.metrics" in parsed.loops
    assert "results.data.critical" in parsed.loops
    assert "results.data.warning" in parsed.loops
    assert "results.data.normal" in parsed.loops
    assert "summary.total" in parsed.placeholders
    assert any(item.endswith("target.ip") for item in parsed.placeholders)
    assert any(item.endswith("target.operating_system") for item in parsed.placeholders)
    assert any("metric_name" in item for item in parsed.placeholders)
    template_text = _rendered_text(fmt, content)
    assert "{{" in template_text or "{%" in template_text
    assert "{d." not in template_text
    assert "Windows" not in template_text
    assert "共 {{ summary.total }} 台主机" in template_text
    assert "主机健康巡检报告" in template_text
    assert "严重问题" in template_text
    assert "作业平台健康巡检报告" not in template_text
    text = _rendered_text(fmt, rendered)
    assert "10.10.90.120" in text
    assert "CPU Total" in text
    assert "Ethernet0" in text
    assert "{{" not in text


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_builtin_health_report_template_renders_multi_host_multi_os(fmt):
    template_path = TEMPLATE_DIR / f"health-inspection-example.{fmt}"
    content = template_path.read_bytes()

    rendered = render_report(content, fmt, MULTI_HOST_REPORT_DATA)
    text = _rendered_text(fmt, rendered)

    assert "10.10.90.120" in text
    assert "192.168.64.5" in text
    assert "windows" in text
    assert "linux" in text
    assert "job-lab-linux-yum" in text
    assert "/mnt/lima-cidata" in text
    assert "共 2 台主机" in text
    assert "{{" not in text
